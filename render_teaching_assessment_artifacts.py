"""Portable artifact authoring from pack.json (python-pptx, python-docx, ReportLab).
Generation does not need Microsoft Office. QA renderers are separate.
"""
import json
from pathlib import Path
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Flowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT=Path(__file__).resolve().parent/'output/teaching_assessment_demo'
FONT='Aptos' # native Office may substitute Calibri; use Arial for broad compatibility below
FONT='Arial'
NAVY='183348'; TEAL='087E8B'; DARK='243B4B'; MUTED='597080'

def textbox(s,x,y,w,h,text,size=22,color=DARK,bold=False):
    shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(12)
    return shape

def native_table(slide,rows,x,y,w,h):
    tbl=slide.shapes.add_table(len(rows),len(rows[0]), Inches(x),Inches(y), Inches(w), Inches(h)).table
    for r,row in enumerate(rows):
        for c,val in enumerate(row):
            cell=tbl.cell(r,c);cell.text=val
            cell.margin_left=Inches(.14);cell.margin_right=Inches(.12)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if r==0 else ('F0F5F7' if r%2 else 'FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(18);p.font.bold=r==0;p.font.color.rgb=RGBColor.from_string('FFFFFF' if r==0 else DARK)
    return tbl

def make_ppt(core):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='Statistics: Exploring and Summarising Data'
    prs.core_properties.subject='Teacher-generated 9MA0 practice; not official Pearson material'
    prs.core_properties.author='AI School Academic OS teaching demo'
    for s in core['slides']:
        sl=prs.slides.add_slide(prs.slide_layouts[6]); n=s['slide_id']
        if n==1:
            sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(NAVY)
            textbox(sl,.8,.7,11, .45,'A LEVEL MATHEMATICS   /   9MA0   /   STATISTICS',17,'A6DBDF',True)
            textbox(sl,.8,1.75,11.5,2.2,'Exploring and\nsummarising data',44,'FFFFFF',True)
            textbox(sl,.8,4.5,11.4,1.4,'Pearson Edexcel qualification context\nTeacher-generated teaching and practice material\nNot official Pearson material',21,'FFFFFF')
        else:
            textbox(sl,.6,.28,12.1,.28,'STATISTICS   /   EXPLORING AND SUMMARISING DATA',12,TEAL,True)
            textbox(sl,.6,.87,12.1,.8,s['title'],32,NAVY,True)
            if s['table']:
                textbox(sl,.65,1.95,6.0,3.65,'\n'.join(s['body']),21)
                native_table(sl,s['table'],7.1,2.02,5.55, min(2.9,.60*len(s['table'])))
            elif s['chart']:
                textbox(sl,.65,1.95,6,3.6,'\n'.join(s['body']),23)
                cd=CategoryChartData();cd.categories=s['chart']['categories']
                for se in s['chart']['series']:cd.add_series(se['name'],se['values'])
                ch=sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7),Inches(1.95), Inches(5.5), Inches(3.5),cd).chart
                ch.has_legend=False;ch.has_title=True;ch.chart_title.text_frame.text='Standard deviation (ml)'
                ch.value_axis.minimum_scale=0;ch.value_axis.maximum_scale=8
                ch.value_axis.tick_labels.font.size=Pt(15);ch.category_axis.tick_labels.font.size=Pt(16)
                ch.plots[0].has_data_labels=True;ch.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
                ch.series[0].format.fill.solid();ch.series[0].format.fill.fore_color.rgb=RGBColor.from_string(TEAL)
                # python-pptx's stock axis IDs can be signed. OOXML requires unsigned
                # integers; normalise both definitions and cross-axis references.
                for el in ch._chartSpace.iter():
                    if el.tag.endswith('}axId') or el.tag.endswith('}crossAx'):
                        el.set('val',str(int(el.get('val')) % (2**32)))
            else:
                textbox(sl,.65,1.95,11.85,3.65,'\n'.join(s['body']),25)
            textbox(sl,.65,5.83,11.9,.27,'STUDENT CHECK',12,TEAL,True)
            textbox(sl,.65,6.22,11.9,.65,s['student_check']['prompt'],19,NAVY)
        textbox(sl,.65,7.12,11.8,.2,f'Teacher-generated • Not official Pearson material   |   Illustrative numerical data   |   {n:02d}',10,'A6DBDF' if n==1 else MUTED)
        notes=s['speaker_notes']+'\n\nStudent check answer: '+s['student_check']['answer']
        notes+='\n\n'+core['material_notice']+'\n'+s['data_notice']
        if s['worked_example_id']:notes+='\nOriginal worked example '+s['worked_example_id']
        notes+='\nCompetency references: '+', '.join(s['competency_ids'])
        sl.notes_slide.notes_text_frame.text=notes
    prs.save(OUT/'Edexcel_9MA0_Statistics_Exploring_Data.pptx')

INSTRUCTIONS=[
 '48 marks. Suggested time: 55 minutes. Answer all eight questions. A calculator may be used.',
 'Show your working. Give non-exact numerical answers to 3 significant figures unless stated otherwise. Include units.',
 'For descriptive variance and standard deviation use divisor n, not n - 1. Quartile/interpolation rules are given in the relevant questions.',
 'All questions are original teacher-generated practice. All numerical datasets are illustrative and invented. They are not Pearson Large Data Set observations.',
 'This is a focused topic assessment, not a full 9MA0 examination or a prediction of official topic weightings.'
]
POLICY='Each listed marking point is worth 1 mark. M = method, A = accuracy, B = independent statement. A normally depends on the preceding M; B is independent. Credit equivalent correct methods. Apply follow-through only where stated. Do not award duplicate credit for one statement.'

def doc_table(doc,rows,widths=None):
    tab=doc.add_table(rows=0, cols=len(rows[0]));tab.autofit=False
    if widths:
        for c,w in zip(tab.columns,widths):c.width=DI(w)
    for i,row in enumerate(rows):
        cells=tab.add_row().cells
        if widths:
            for c,w in zip(cells,widths):c.width=DI(w)
        for cell,text in zip(cells,row):
            cell.text=text
            props=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF2' if i==0 else 'FFFFFF');props.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            props.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=DP(6);p.paragraph_format.space_before=DP(5)
                for r in p.runs:r.font.size=DP(10);r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=DP(3)
    return tab

def new_doc(title,scheme):
    d=Document();sec=d.sections[0];sec.page_width=DI(8.2677);sec.page_height=DI(11.6929)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec.top_margin=DI(.7);sec.bottom_margin=DI(.7);sec.left_margin=DI(.75);sec.right_margin=DI(.75)
    for name in ['Normal','Title','Heading 1','Heading 2','Subtitle']:
        st=d.styles[name];st.font.name=FONT;st.font.color.rgb=DC(0,0,0)
    d.styles['Normal'].font.size=DP(11);d.styles['Normal'].paragraph_format.space_after=DP(8)
    d.styles['Title'].font.size=DP(25);d.styles['Heading 1'].font.size=DP(17)
    header=sec.header.paragraphs[0];header.text='9MA0 Mathematics  •  Exploring and Summarising Data';header.runs[0].font.size=DP(9)
    foot=sec.footer.paragraphs[0];foot.text='Teacher-generated practice. Not official Pearson material.     Page '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for r in foot.runs:r.font.size=DP(8)
    d.add_paragraph(title,'Title');d.add_paragraph('Pearson Edexcel A Level Mathematics (9MA0)\nStatistics / Applied Mathematics')
    return d

def make_docs(core):
    for scheme in [False,True]:
        name='Mark_Scheme' if scheme else 'Assessment'
        d=new_doc('Statistics Practice Mark Scheme' if scheme else 'Statistics Practice Assessment',scheme)
        if scheme:
            d.add_paragraph(core['mark_scheme_notice'])
            d.add_paragraph('48 marks across 8 questions and 21 question parts. This scheme accompanies the original practice assessment.')
            d.add_paragraph(POLICY)
            d.add_paragraph('Numerical answer checks: Q2 mean 11; median 8. Q3 estimated mean 15; 75th percentile 20. Q4 range 26; Q1 7.5; Q3 16.5; IQR 9; upper fence 30. Q6 variance 9; SD 3. Q7 available-day mean 1.2. Q8 revised mean 14.')
            d.add_paragraph('The stated quartile convention and missing/trace policy are deliberate teaching choices. Review demand and timing for the class before use. No official grade boundaries are supplied.')
        else:
            d.add_paragraph('Name: __________________________________    Class: ______________')
            d.add_paragraph('Date: ___________________')
            for t in INSTRUCTIONS:d.add_paragraph(t)
            d.add_paragraph('Useful formulae','Heading 1')
            d.add_paragraph('Mean = sum x / n\nGrouped estimated mean = sum(f x midpoint) / sum f\nVariance = sum x² / n - mean²\nStandard deviation = square root of variance')
            d.add_paragraph('Use the working space provided. If you need additional paper, label each answer clearly.')
        for q in core['questions']:
            d.add_page_break();total=sum(p['marks'] for p in q['parts'])
            d.add_paragraph(q['question_id']+'  '+q['title']+f'  ({total} marks)','Heading 1')
            if not scheme:
                d.add_paragraph(q['context'])
                if q['data']:doc_table(d,q['data'])
                for p in q['parts']:
                    d.add_paragraph(p['part_id']+'  '+p['prompt']+f"  [{p['marks']} marks]")
                    # About 5–6 mm per line, and more space for higher-mark parts.
                    for _ in range(3 if p['marks']<=2 else 4):
                        line=d.add_paragraph('________________________________________________________________________________')
                        line.paragraph_format.space_after=DP(5);line.paragraph_format.space_before=DP(0)
                        for r in line.runs:r.font.size=DP(8);r.font.color.rgb=DC(180,180,180)
            else:
                rows=[['Part','Answer and method','Marks / allocation']]
                for p in q['parts']:rows.append([p['part_id'],p['answer'],str(p['marks'])+'\n'+'\n'.join(p['marking_points'])])
                doc_table(d,rows,[.57,3.57,2.60])
                for p in q['parts']:
                    d.add_paragraph(p['part_id']+' notes: '+p['acceptable_alternatives']+' '+p['interpretation_requirements'])
        d.save(OUT/f'Edexcel_9MA0_Statistics_{name}.docx')

class WorkingLines(Flowable):
    def __init__(self,count):super().__init__();self.count=count;self.width=480;self.height=count*19+7
    def draw(self):
        self.canv.setStrokeColor(colors.HexColor('#CBD3D9'));self.canv.setLineWidth(.4)
        for i in range(self.count):self.canv.line(0,self.height-19*(i+1),480,self.height-19*(i+1))

def make_pdfs(core):
    fonts=Path('C:/Windows/Fonts')
    pdfmetrics.registerFont(TTFont('TeacherArial',str(fonts/'arial.ttf')))
    pdfmetrics.registerFont(TTFont('TeacherArialBold',str(fonts/'arialbd.ttf')))
    styles=getSampleStyleSheet()
    for name in ['Normal','Title','Heading1','Heading2']:
        styles[name].fontName='TeacherArialBold' if name!='Normal' else 'TeacherArial'
        styles[name].textColor=colors.black
    styles['Normal'].fontSize=11;styles['Normal'].leading=16;styles['Normal'].spaceAfter=9
    styles['Title'].fontSize=25;styles['Title'].leading=30;styles['Title'].alignment=0
    styles['Heading1'].fontSize=17;styles['Heading1'].leading=22
    small=ParagraphStyle('small',parent=styles['Normal'],fontSize=9,leading=12,spaceAfter=3)
    def para(t,style='Normal'):return Paragraph(escape(t).replace('\n','<br/>'),styles[style])
    def table(rows,scheme=False):
        widths=[40,258,185] if scheme else [483/len(rows[0])]*len(rows[0])
        t=Table([[Paragraph(escape(c).replace('\n','<br/>'),small) for c in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EDF2')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]));return t
    def footer(canvas,doc):
        canvas.setFont('TeacherArial',8);canvas.setFillColor(colors.HexColor('#4B5961'))
        canvas.drawString(56,815,'9MA0 Mathematics   /   Exploring and Summarising Data')
        canvas.drawString(56,28,'Teacher-generated practice. Not official Pearson material.')
        canvas.drawRightString(539,28,f'Page {doc.page}')
    for scheme in [False,True]:
        name='Mark_Scheme' if scheme else 'Assessment';story=[]
        story += [para('Statistics Practice Mark Scheme' if scheme else 'Statistics Practice Assessment','Title'),Spacer(1,12),para('Pearson Edexcel A Level Mathematics (9MA0)\nStatistics / Applied Mathematics'),Spacer(1,10)]
        if scheme:
            story += [para(core['mark_scheme_notice']),para('48 marks. 8 questions. 21 question parts.'),para(POLICY),para('Mark all responses against the stated data and conventions. No official grade boundaries are supplied. Each question is original practice, and all numerical data are illustrative.'),para('Human review before use','Heading1'),para('Check calculator conventions, the raw-data quartile rule, local marking preferences and timing for your class. Predicted demand is a teacher estimate, not a validated psychometric scale.')]
        else:
            story += [para('Name: __________________________________    Class: __________'),para('Date: ___________________'),Spacer(1,10)]
            story += [para(t) for t in INSTRUCTIONS]
            story += [para('Useful formulae','Heading1'),para('Mean = sum x / n\nGrouped estimated mean = sum(f × midpoint) / sum f\nVariance = sum x² / n − mean²\nStandard deviation = square root of variance'),Spacer(1,12),para('Use the working space provided. Label any additional sheets clearly.')]
        for q in core['questions']:
            story += [PageBreak(),para(q['question_id']+'  '+q['title']+f"  ({sum(p['marks'] for p in q['parts'])} marks)",'Heading1')]
            if scheme:
                rows=[['Part','Answer and method','Marks / allocation']]
                rows += [[p['part_id'],p['answer'],str(p['marks'])+'\n'+'\n'.join(p['marking_points'])] for p in q['parts']]
                story += [table(rows,True),Spacer(1,15)]
                story += [para(p['part_id']+' notes: '+p['acceptable_alternatives']+' '+p['interpretation_requirements']) for p in q['parts']]
            else:
                story += [para(q['context'])]
                if q['data']:story += [table(q['data']),Spacer(1,12)]
                for p in q['parts']:
                    story += [para(p['part_id']+'  '+p['prompt']+f"  [{p['marks']} marks]"),WorkingLines(3 if p['marks']<=2 else 4)]
        doc=SimpleDocTemplate(str(OUT/f'Edexcel_9MA0_Statistics_{name}.pdf'),pagesize=A4,leftMargin=56,rightMargin=56,topMargin=52,bottomMargin=48,title='Statistics practice '+name,author='AI School Academic OS teaching demo')
        doc.build(story,onFirstPage=footer,onLaterPages=footer)

def main():
    core=json.loads((OUT/'pack.json').read_text(encoding='utf-8'))
    make_ppt(core);make_docs(core);make_pdfs(core)
    print('Created 1 PPTX, 2 DOCX, 2 PDF artifacts.')

if __name__=='__main__':main()
